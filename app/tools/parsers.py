"""PDF and DOCX parsing with page preservation and OCR fallback."""
from __future__ import annotations

import io
from dataclasses import dataclass, field
from typing import Any

from app.core.errors import DocumentProcessingError, OCRUnavailableError
from app.core.logging import get_logger
from app.tools.normalize import normalize_page_text, strip_repeated_lines
from app.tools.ocr import OcrEngine, UnavailableOcr

log = get_logger(__name__)
MIN_TEXT_CHARS_PER_PAGE = 40
LOGICAL_PAGE_CHARS = 3000


@dataclass
class ParsedPage:
    number: int
    text: str
    ocr_used: bool = False
    ocr_confidence: float | None = None
    needs_ocr: bool = False
    tables: int = 0


@dataclass
class PageSpan:
    page: int
    start: int
    end: int


@dataclass
class ParsedDocument:
    pages: list[ParsedPage]
    page_basis: str = "physical"  # physical | logical
    warnings: list[str] = field(default_factory=list)
    metadata: dict[str, Any] = field(default_factory=dict)

    @property
    def is_scanned(self) -> bool:
        return any(p.ocr_used or p.needs_ocr for p in self.pages)

    @property
    def ocr_used(self) -> bool:
        return any(p.ocr_used for p in self.pages)

    @property
    def ocr_confidence(self) -> float | None:
        confs = [p.ocr_confidence for p in self.pages if p.ocr_used and p.ocr_confidence is not None]
        return sum(confs) / len(confs) if confs else None

    def assemble(self) -> tuple[str, list[PageSpan]]:
        """Concatenate pages ("\n\n" separated) and return per-page character spans."""
        parts: list[str] = []
        spans: list[PageSpan] = []
        pos = 0
        for p in self.pages:
            if parts:
                pos += 2
            spans.append(PageSpan(p.number, pos, pos + len(p.text)))
            parts.append(p.text)
            pos += len(p.text)
        return "\n\n".join(parts), spans


def page_for_offset(spans: list[PageSpan], offset: int) -> int:
    lo, hi = 0, len(spans) - 1
    if not spans:
        return 1
    while lo <= hi:
        mid = (lo + hi) // 2
        s = spans[mid]
        if offset < s.start:
            hi = mid - 1
        elif offset > s.end:
            lo = mid + 1
        else:
            return s.page
    return spans[min(max(lo, 0), len(spans) - 1)].page


# ---------------------------------------------------------------------------------------------
def parse_pdf(data: bytes, *, ocr: OcrEngine | None = None, max_pages: int = 600, ocr_dpi: int = 200) -> ParsedDocument:
    import pymupdf as fitz  # PyMuPDF

    ocr = ocr or UnavailableOcr()
    try:
        doc = fitz.open(stream=data, filetype="pdf")
    except Exception as exc:  # noqa: BLE001
        raise DocumentProcessingError(f"cannot open pdf: {exc}", user_message="The PDF could not be opened. It may be corrupt.") from exc
    with doc:
        if doc.needs_pass:
            raise DocumentProcessingError("encrypted pdf", user_message="The PDF is password-protected. Remove the password and upload it again.")
        if doc.page_count == 0:
            raise DocumentProcessingError("no pages", user_message="The PDF contains no pages.")
        if doc.page_count > max_pages:
            raise DocumentProcessingError(f"{doc.page_count} pages", user_message=f"The PDF has {doc.page_count} pages; the limit is {max_pages}.")
        pages: list[ParsedPage] = []
        warnings: list[str] = []
        for index in range(doc.page_count):
            page = doc.load_page(index)
            try:
                raw = page.get_text("text", sort=True)
            except Exception as exc:  # noqa: BLE001
                log.warning("text extraction failed on page %d: %s", index + 1, type(exc).__name__)
                raw = ""
            tables = 0
            try:
                tables = len(page.find_tables().tables)
            except Exception:  # noqa: BLE001 - table detection is best-effort
                tables = 0
            pp = ParsedPage(number=index + 1, text=raw, tables=tables)
            if len(raw.strip()) < MIN_TEXT_CHARS_PER_PAGE:
                pp.needs_ocr = True
                if ocr.is_available():
                    try:
                        png = page.get_pixmap(dpi=ocr_dpi).tobytes("png")
                        res = ocr.ocr_image(png)
                        pp.text, pp.ocr_confidence, pp.ocr_used, pp.needs_ocr = res.text, res.confidence, True, False
                    except OCRUnavailableError:
                        raise
                    except Exception as exc:  # noqa: BLE001
                        warnings.append(f"OCR failed on page {index + 1}: {type(exc).__name__}")
                else:
                    warnings.append(f"Page {index + 1} has no extractable text and OCR is unavailable.")
            pages.append(pp)
        meta = {k: v for k, v in (doc.metadata or {}).items() if v}
    return _finalize(pages, "physical", warnings, meta, ocr)


def _finalize(pages: list[ParsedPage], basis: str, warnings: list[str], meta: dict[str, Any], ocr: OcrEngine | None = None) -> ParsedDocument:
    usable = [p for p in pages if p.text.strip()]
    if not usable:
        if any(p.needs_ocr for p in pages):
            raise OCRUnavailableError(
                "scanned document without OCR",
                user_message="This document is scanned (no text layer) and OCR is not available. Install Tesseract OCR and set TESSERACT_CMD, then upload again.",
            )
        raise DocumentProcessingError("no text", user_message="No readable text was found in this document.")
    cleaned = strip_repeated_lines([normalize_page_text(p.text) for p in pages])
    for p, text in zip(pages, cleaned):
        p.text = text
    return ParsedDocument(pages=pages, page_basis=basis, warnings=warnings, metadata=meta)


# ---------------------------------------------------------------------------------------------
_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _paragraph_text(p_el: Any) -> tuple[str, bool, bool]:
    """Text of a w:p including tracked insertions (excluding deletions).

    Returns ``(text, break_before, break_after)``; a page marker seen before any text starts a new
    page, one seen after text ends the current page.
    """
    parts: list[str] = []
    before = after = False
    for el in p_el.iter():
        tag = el.tag
        marker = False
        if tag == _W + "t":
            parts.append(el.text or "")
        elif tag == _W + "tab":
            parts.append("\t")
        elif tag == _W + "br":
            if el.get(_W + "type") == "page":
                marker = True
            else:
                parts.append("\n")
        elif tag == _W + "lastRenderedPageBreak":
            marker = True
        if marker:
            if "".join(parts).strip():
                after = True
            else:
                before = True
    return "".join(parts), before, after


def parse_docx(data: bytes) -> ParsedDocument:
    from docx import Document as DocxDocument

    try:
        doc = DocxDocument(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        raise DocumentProcessingError(f"cannot open docx: {exc}", user_message="The DOCX file could not be opened. It may be corrupt.") from exc
    blocks: list[tuple[str, bool, bool]] = []
    table_count = 0
    for child in doc.element.body.iterchildren():
        if child.tag == _W + "p":
            blocks.append(_paragraph_text(child))
        elif child.tag == _W + "tbl":
            table_count += 1
            for tr in child.iter(_W + "tr"):
                cells = []
                for tc in tr.iter(_W + "tc"):
                    cells.append(" ".join(_paragraph_text(p)[0].strip() for p in tc.iter(_W + "p")).strip())
                blocks.append((" | ".join(c for c in cells), False, False))
    explicit = any(b or a for _, b, a in blocks)
    pages_text: list[list[str]] = [[]]
    size = 0
    for text, brk_before, brk_after in blocks:
        if explicit and brk_before and any(x.strip() for x in pages_text[-1]):
            pages_text.append([])
            size = 0
        pages_text[-1].append(text)
        size += len(text) + 1
        if explicit and brk_after:
            pages_text.append([])
            size = 0
        elif not explicit and size >= LOGICAL_PAGE_CHARS:
            pages_text.append([])
            size = 0
    pages = [ParsedPage(number=0, text="\n".join(t)) for t in pages_text if any(x.strip() for x in t)]
    for i, p in enumerate(pages):
        p.number = i + 1
    if pages:
        pages[0].tables = table_count
    warnings: list[str] = []
    if not explicit:
        warnings.append("DOCX has no physical page information; page numbers are logical (about 3,000 characters each).")
    core = doc.core_properties
    meta = {k: v for k, v in {"title": core.title, "author": core.author}.items() if v}
    return _finalize(pages, "physical" if explicit else "logical", warnings, meta)


def parse_document(filename: str, data: bytes, *, ocr: OcrEngine | None = None, max_pages: int = 600, ocr_dpi: int = 200) -> ParsedDocument:
    lower = filename.lower()
    if lower.endswith(".pdf"):
        return parse_pdf(data, ocr=ocr, max_pages=max_pages, ocr_dpi=ocr_dpi)
    if lower.endswith(".docx"):
        return parse_docx(data)
    raise DocumentProcessingError(f"unsupported extension: {filename}", user_message="Unsupported file type.")

