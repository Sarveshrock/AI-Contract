"""Generate sample PDF/DOCX files under data/samples from app/demo/sample_texts.py.

Usage: python scripts/generate_samples.py
Requires reportlab (dev dependency) and PyMuPDF.
"""
from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from app.demo.sample_texts import SAMPLES  # noqa: E402

OUT = ROOT / "data" / "samples"


def write_pdf(name: str, text: str, path: Path) -> None:
    from reportlab.lib.pagesizes import LETTER
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
    from xml.sax.saxutils import escape

    styles = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=styles["BodyText"], fontName="Helvetica", fontSize=10.5, leading=14)
    head = ParagraphStyle("head", parent=body, fontName="Helvetica-Bold", fontSize=11, spaceBefore=8, keepWithNext=1)
    title = ParagraphStyle("title", parent=head, fontSize=15, alignment=1, spaceAfter=10)

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.drawCentredString(LETTER[0] / 2, 0.5 * inch, f"{name} - Confidential - Page {doc.page}")
        canvas.restoreState()

    story = []
    first = True
    for block in text.strip().split("\n"):
        if not block.strip():
            story.append(Spacer(1, 4))
            continue
        style = title if first else (head if block.isupper() or (block[:2].strip(". ").isdigit() and block.split(" ", 1)[1].isupper()) else body)
        story.append(Paragraph(escape(block), style))
        first = False
    SimpleDocTemplate(str(path), pagesize=LETTER, leftMargin=inch, rightMargin=inch, topMargin=inch, bottomMargin=0.9 * inch).build(story, onFirstPage=footer, onLaterPages=footer)


def write_docx(text: str, path: Path) -> None:
    from docx import Document

    doc = Document()
    for i, line in enumerate(text.strip().split("\n")):
        if not line.strip():
            continue
        if i == 0:
            doc.add_heading(line, level=0)
        elif line.isupper() or (line[:2].strip(". ").isdigit() and line.split(" ", 1)[1].isupper()):
            doc.add_heading(line, level=2)
        else:
            doc.add_paragraph(line)
    doc.save(str(path))


def write_scanned_pdf(source_pdf: Path, path: Path) -> None:
    """Rasterise the first page of a PDF into an image-only PDF (no text layer) to exercise OCR."""
    import pymupdf as fitz

    with fitz.open(str(source_pdf)) as src:
        png = src[0].get_pixmap(dpi=130).tobytes("jpeg")
    out = fitz.open()
    page = out.new_page(width=612, height=792)
    page.insert_image(page.rect, stream=png)
    out.save(str(path))
    out.close()


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    for name, text in SAMPLES.items():
        if name == "msa_v2_restated":
            write_docx(text, OUT / f"{name}.docx")
        else:
            write_pdf(name, text, OUT / f"{name}.pdf")
        print("wrote", name)
    write_scanned_pdf(OUT / "mutual_nda.pdf", OUT / "scanned_nda_image_only.pdf")
    print("wrote scanned_nda_image_only")


if __name__ == "__main__":
    main()
